from io import BytesIO
from pathlib import Path

import logging

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils.text import slugify
from openpyxl import load_workbook

from .models import Acta

TEMPLATE_PATH = Path(settings.BASE_DIR) / "templates" / "actas" / "acta_template.xlsx"
logger = logging.getLogger(__name__)


def _value(values, section, label, default="N.A"):
    return values.get((section, label), default) or default


def _set(ws, cell, value):
    for merged_range in ws.merged_cells.ranges:
        if cell in merged_range:
            cell = merged_range.start_cell.coordinate
            break
    ws[cell] = value if value not in (None, "") else "N.A"


def _peripheral(values, name):
    raw = _value(values, "perifericos", name, "N.A")
    parts = [part.strip() for part in str(raw).split("/")]
    parts += ["N.A"] * (4 - len(parts))
    return parts[:4]


def _field_values(acta):
    return {
        (item.definition.section, item.definition.label): item.value
        for item in acta.field_values.select_related("definition").all()
    }


def render_acta_xlsx(acta: Acta) -> BytesIO:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"No existe la plantilla Excel: {TEMPLATE_PATH}")

    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["FORMATO"]
    values = _field_values(acta)
    equipment = lambda label, default="N.A": _value(values, "equipo", label, default)

    _set(sheet, "B4", acta.act_date)
    _set(sheet, "F4", acta.site.name if acta.site else "N.A")
    _set(sheet, "B5", "SI" if acta.act_type.lower() == "entrega" else "N.A")
    _set(sheet, "G5", "SI" if acta.act_type.lower() == "cambio" else "N.A")
    _set(sheet, "A6", f"N° CASO GLPI: {acta.glpi_case_number}")
    _set(sheet, "E6", f"PROCESO/PORTAFOLIO: {acta.portfolio or 'N.A'}")
    _set(sheet, "A7", f"NOMBRE COMPLETO: {_value(values, 'usuario', 'Nombre completo')}")
    _set(sheet, "E7", f"CÉDULA: {_value(values, 'usuario', 'Cédula')}")

    _set(sheet, "B10", _value(values, "sistemas", "Tipo de cuenta"))
    _set(sheet, "C10", "SI" if _value(values, "sistemas", "Tipo de cuenta", "") else "N.A")
    _set(sheet, "E10", "Correo")
    _set(sheet, "G10", "SI" if _value(values, "sistemas", "Cuenta de correo", "") != "N.A" else "N.A")
    _set(sheet, "A11", f"Windows: {_value(values, 'sistemas', 'Usuario de Windows')}")
    _set(sheet, "E11", f"Cuenta de Correo: {_value(values, 'sistemas', 'Cuenta de correo')}")
    _set(sheet, "A12", f"Acceso a HelpDesk (GLPI): {'SI' if _value(values, 'sistemas', 'Acceso a HelpDesk GLPI', False) else 'N.A'}")
    _set(sheet, "E12", f"Novedades AIO: {'SI' if _value(values, 'sistemas', 'Novedades AIO', False) else 'N.A'}")
    _set(sheet, "A13", f"Impresora: {'SI' if _value(values, 'sistemas', 'Impresora', False) else 'N.A'}")
    _set(sheet, "E13", f"Control de Acceso: {'SI' if _value(values, 'sistemas', 'Control de acceso', False) else 'N.A'}")
    _set(sheet, "A14", f"Recurso Compartido: {_value(values, 'sistemas', 'Recurso compartido')}")
    _set(sheet, "E14", f"Licencia Office365: {'SI' if _value(values, 'sistemas', 'Licencia Office 365', False) else 'N.A'}")
    _set(sheet, "A15", f"Extensión: {_value(values, 'sistemas', 'Extensión')}")
    _set(sheet, "E15", f"Perfil de Navegación: {_value(values, 'sistemas', 'Perfil de navegación')}")
    _set(sheet, "A16", f"Power BI: {'SI' if _value(values, 'sistemas', 'Power BI', False) else 'N.A'}")
    _set(sheet, "E16", f"VPN: {'SI' if _value(values, 'sistemas', 'VPN', False) else 'N.A'}")

    _set(sheet, "B19", "SI" if equipment("Tipo de equipo", "").lower() == "desktop" else "N.A")
    _set(sheet, "G19", "SI" if equipment("Tipo de equipo", "").lower() == "portátil" else "N.A")
    _set(sheet, "A20", f"Equipo/Modelo: {equipment('Marca')} {equipment('Modelo')}")
    _set(sheet, "D20", f"Serial: {equipment('Serial')}")
    _set(sheet, "F20", f"Serial Devuelto (Si es cambio): {equipment('Serial devuelto')}")
    _set(sheet, "A21", f"Procesador: {equipment('Procesador')}")
    _set(sheet, "D21", f"Dirección IP: {equipment('Dirección IP')}")
    _set(sheet, "F21", f"Memoria RAM: {equipment('Memoria RAM')}")
    _set(sheet, "A22", f"Disco Duro: {equipment('Disco duro')}")
    _set(sheet, "D22", f"Sistema Operativo: {equipment('Sistema operativo')}")
    _set(sheet, "F22", f"Placa Inventario: {equipment('Placa de inventario')}")

    peripheral_names = [("Diadema", "A25", "C25"), ("CPU", "E25", "G25"), ("Teclado", "A26", "C26"), ("Móvil", "E26", "G26"), ("Mouse", "A27", "C27"), ("SIM", "E27", "G27"), ("Otro", "A28", "C28")]
    for name, label_cell, status_cell in peripheral_names:
        present = _value(values, "perifericos", name, "N.A") != "N.A"
        _set(sheet, status_cell, "SI" if present else "N.A")
    peripheral = next((_peripheral(values, name) for name in ("Diadema", "Teclado", "Mouse", "CPU", "Móvil", "SIM", "Otro") if _value(values, "perifericos", name, "N.A") != "N.A"), ["N.A"] * 4)
    _set(sheet, "A29", f"Marca/Modelo: {peripheral[0]} {peripheral[1]}")
    _set(sheet, "E29", f"Placa Inventario GLPI: {peripheral[2]}")
    _set(sheet, "A30", f"Serial: {peripheral[3]}")
    _set(sheet, "E30", f"Serial Devuelto (Si es cambio): {equipment('Serial devuelto')}")

    _set(sheet, "A32", "SI" if _value(values, "monitor", "Marca", "N.A") != "N.A" else "N.A")
    _set(sheet, "A33", f"Marca/Modelo: {_value(values, 'monitor', 'Marca')} {_value(values, 'monitor', 'Modelo')}")
    _set(sheet, "E33", f"Placa Inventario GLPI: {_value(values, 'monitor', 'Placa de inventario')}")
    _set(sheet, "A34", f"Serial: {_value(values, 'monitor', 'Serial')}")
    _set(sheet, "A36", _value(values, "observaciones", "Observaciones", "N.A"))

    signatures = {signature.signature_type: signature for signature in acta.signatures.all()}
    for cell, signature_type in (("B43", "DELIVERY"), ("D43", "REVIEW"), ("F43", "FINAL_APPROVAL"), ("H43", "RECEIVE")):
        signature = signatures.get(signature_type)
        _set(sheet, cell, signature.signer_name if signature else "Pendiente")
    for cell in ("B44", "D44", "F44", "H44"):
        _set(sheet, cell, "Pendiente")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def render_acta_pdf(acta: Acta) -> bytes:
    from django.template.loader import render_to_string
    from weasyprint import HTML

    from .html_views import _acta_document_context

    markup = render_to_string(
        "actas/_acta_document.html",
        {"acta": acta, **_acta_document_context(acta)},
    )
    return HTML(string=markup, base_url=str(settings.BASE_DIR)).write_pdf()


def save_final_acta_pdf(acta: Acta) -> None:
    pdf_bytes = render_acta_pdf(acta)
    filename = f"acta-{acta.pk}-{slugify(acta.glpi_case_number)}.pdf"
    acta.pdf_file.save(filename, ContentFile(pdf_bytes), save=False)


def upload_final_acta_to_glpi(acta: Acta, actor=None) -> bool:
    if not settings.GLPI_UPLOAD_ENABLED or not acta.pdf_file:
        return False

    from glpi.client import GlpiClient
    import requests

    previous_status = acta.status
    try:
        with acta.pdf_file.open("rb") as pdf_file:
            GlpiClient().upload_document(
                acta.glpi_case_number,
                acta.pdf_file.name.rsplit("/", 1)[-1],
                pdf_file,
            )
        acta.status = "GLPI_UPLOADED"
        acta.glpi_upload_status = "UPLOADED"
        acta.glpi_uploaded_at = acta.completed_at
        acta.save(update_fields=["status", "glpi_upload_status", "glpi_uploaded_at", "updated_at"])
        from audit.models import ActaEvent
        ActaEvent.objects.create(
            acta=acta,
            event_type="GLPI",
            action="DOCUMENT_UPLOADED",
            actor=actor,
            from_status=previous_status,
            to_status=acta.status,
        )
        return True
    except (requests.RequestException, RuntimeError, ValueError, OSError):
        logger.exception("No se pudo cargar el PDF final a GLPI para acta=%s", acta.public_id)
        acta.status = "GLPI_UPLOAD_FAILED"
        acta.glpi_upload_status = "FAILED"
        acta.save(update_fields=["status", "glpi_upload_status", "updated_at"])
        from audit.models import ActaEvent
        ActaEvent.objects.create(
            acta=acta,
            event_type="GLPI",
            action="DOCUMENT_UPLOAD_FAILED",
            actor=actor,
            from_status=previous_status,
            to_status=acta.status,
        )
        return False
