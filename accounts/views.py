from datetime import date
import logging
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.utils import timezone
from django.conf import settings
from openpyxl import Workbook
from reports.models import PilotFeedback
from reports.services import report_data
from django.contrib.auth import get_user_model
from actas.models import Acta
from notifications.models import Notification
from .permissions import active_role
from .permissions import visible_actas
from django.forms import ImageField
from django.core.validators import FileExtensionValidator
from django import forms

logger = logging.getLogger(__name__)


def csrf_failure(request, reason=""):
    if request.path == "/login/" and request.method == "POST" and reason:
        messages.warning(request, "El formulario de acceso venció. Vuelve a intentarlo con el formulario actualizado.")
        return redirect("login")
    return HttpResponse("La solicitud no superó la verificación de seguridad. Recarga el formulario e inténtalo de nuevo.", status=403)


class SupervisorSignatureForm(forms.Form):
    signature = ImageField(validators=[FileExtensionValidator(["png"])])

    def clean_signature(self):
        upload = self.cleaned_data["signature"]
        if upload.size > 1024 * 1024:
            raise forms.ValidationError("La firma PNG no puede superar 1 MB.")
        if upload.content_type != "image/png":
            raise forms.ValidationError("El archivo debe ser una imagen PNG.")
        header = upload.read(8)
        upload.seek(0)
        if header != b"\x89PNG\r\n\x1a\n":
            raise forms.ValidationError("El contenido del archivo no es PNG.")
        return upload


@login_required
def supervisor_signature(request):
    role = active_role(request.user)
    profile = getattr(request.user, "role_profile", None)
    if role not in {"TECHNICIAN", "SUPERVISOR_ONE", "SUPERVISOR_TWO"} or profile is None:
        return HttpResponse("Solo t\u00e9cnicos y supervisores pueden gestionar su firma digital.", status=403)
    form = SupervisorSignatureForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        profile.digital_signature = form.cleaned_data["signature"]
        profile.save(update_fields=["digital_signature"])
        messages.success(request, "Firma digital PNG actualizada.")
        return redirect("accounts:supervisor-signature")
    return render(request, "accounts/supervisor_signature.html", {"form": form, "profile": profile, "title": "Firma digital del usuario"})


ROLE_LABELS = {
    "ADMIN": "Administrador",
    "TECHNICIAN": "Técnico",
    "SUPERVISOR_ONE": "Supervisor uno",
    "SUPERVISOR_TWO": "Supervisor dos",
    "RECEIVER": "Receptor del activo",
    "AUDITOR": "Auditor",
}


@login_required
def dashboard(request):
    role = active_role(request.user)
    if not role:
        return HttpResponse("La cuenta no tiene un rol activo asignado.", status=403)
    qs = Acta.objects.select_related("site", "created_by", "assigned_technician").order_by("-created_at")

    if role == "TECHNICIAN":
        qs = qs.filter(assigned_technician=request.user)
        template_name = "dashboard_technician.html"
    elif role == "RECEIVER":
        qs = qs.filter(receiver=request.user)
        template_name = "dashboard_technician.html"
    elif role in {"SUPERVISOR_ONE", "SUPERVISOR_TWO"}:
        if role == "SUPERVISOR_ONE":
            qs = qs.filter(supervisor_one=request.user, status__in=["PENDING_SUPERVISOR_ONE", "DELIVERY_SIGNED", "REJECTED_BY_SUPERVISOR_ONE"])
        else:
            qs = qs.filter(supervisor_two=request.user, status__in=["PENDING_SUPERVISOR_TWO", "REJECTED_BY_SUPERVISOR_TWO"])
        template_name = "dashboard_supervisor.html"
    elif role == "ADMIN":
        qs = qs.all()
        template_name = "dashboard_admin.html"
    elif role == "AUDITOR":
        qs = qs.all()
        template_name = "dashboard_auditor.html"
    else:
        return HttpResponse("Rol no reconocido.", status=403)

    stats = {
        "total": qs.count(),
        "draft": qs.filter(status="DRAFT").count(),
        "pending": qs.filter(status__in=["PENDING_SUPERVISOR_ONE", "PENDING_SUPERVISOR_TWO"]).count(),
        "completed": qs.filter(status="COMPLETED").count(),
    }

    latest = qs[:6]
    notifications = []
    for acta in qs[:4]:
        if acta.status == "COMPLETED":
            title = "Acta completada"
            description = f"La acta {acta.glpi_case_number} fue cerrada correctamente."
        elif "REJECTED" in acta.status:
            title = "Acta rechazada"
            description = f"La acta {acta.glpi_case_number} requiere revisión por: {acta.rejection_reason or 'motivo no indicado'}."
        else:
            title = "Acta pendiente"
            description = f"La acta {acta.glpi_case_number} está esperando aprobación."

        notifications.append({
            "title": title,
            "description": description,
            "time": acta.updated_at.strftime("%d/%m/%Y %H:%M"),
        })

    if not notifications:
        notifications = [{
            "title": "Sin actividad",
            "description": "Todavía no hay actas registradas para esta vista.",
            "time": "Ahora",
        }]

    context = {
        "user": request.user,
        "role": role,
        "role_label": ROLE_LABELS.get(role, "Usuario"),
        "stats": stats,
        "latest_actas": latest,
        "notifications": notifications,
        "title": "Dashboard",
        "users": get_user_model().objects.select_related("role_profile").order_by("username"),
    }
    return render(request, f"dashboard/{template_name}", context)


@login_required
def notifications(request):
    notification_rows = Notification.objects.filter(recipient=request.user).select_related("acta")[:20]
    notifications = [{
        "title": item.title,
        "description": item.message,
        "time": item.created_at.strftime("%d/%m/%Y %H:%M"),
        "type": item.level.lower(),
    } for item in notification_rows]

    if not notifications:
        notifications = [{
            "title": "Sin notificaciones",
            "description": "No hay actas registradas todavía.",
            "time": "Ahora",
            "type": "success",
        }]

    role = active_role(request.user)
    return render(request, "notifications.html", {"notifications": notifications, "role_label": ROLE_LABELS.get(role, "Usuario"), "title": "Notificaciones"})


@login_required
def reports(request):
    role = active_role(request.user)
    if role not in {"ADMIN", "SUPERVISOR_ONE", "SUPERVISOR_TWO", "AUDITOR"}:
        return HttpResponse("No tienes permiso para consultar reportes.", status=403)
    start_date = _parse_date(request.GET.get("desde"))
    end_date = _parse_date(request.GET.get("hasta"))
    data = report_data(start_date, end_date, visible_actas(Acta.objects.all(), request.user))
    context = {
        "title": "Reportes",
        "report": data,
        "desde": start_date.isoformat() if start_date else "",
        "hasta": end_date.isoformat() if end_date else "",
        "role_label": ROLE_LABELS.get(role, "Usuario"),
    }
    return render(request, "reports.html", context)


def _parse_date(value):
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


@login_required
def reports_export(request, export_format):
    if active_role(request.user) not in {"ADMIN", "SUPERVISOR_ONE", "SUPERVISOR_TWO", "AUDITOR"}:
        return HttpResponse("No tienes permiso para exportar reportes.", status=403)
    data = report_data(_parse_date(request.GET.get("desde")), _parse_date(request.GET.get("hasta")), visible_actas(Acta.objects.all(), request.user))
    if export_format == "xlsx":
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Resumen"
        summary.append(["Métrica", "Valor"])
        summary.append(["Actas en periodo", data["actas_total"]])
        summary.append(["Actas con firmas", data["actas_firmadas"]])
        summary.append([])
        summary.append(["Actas completadas por mes", "Cantidad"])
        for period, count in data["signed_by_period"]:
            summary.append([period, count])
        equipment = workbook.create_sheet("Equipos")
        equipment.append(["Marca", "Modelo", "Cantidad"])
        for row in data["equipment"]:
            equipment.append([row["brand"], row["model"], row["count"]])
        stages = workbook.create_sheet("Aprobaciones")
        stages.append(["Etapa", "Responsable", "Promedio horas", "Actas"])
        for row in data["supervisor_averages"]:
            stages.append([row["stage"], row["user"], row["hours"], row["count"]])
        pending = workbook.create_sheet("Pendientes")
        pending.append(["Caso GLPI", "Etapa", "Responsable", "Días pendiente", "Estado"])
        for row in data["pending"]:
            pending.append([row["case"], row["stage"], row["owner"], row["days"], row["status"]])
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="reportes-actas.xlsx"'
        workbook.save(response)
        return response
    if export_format == "pdf":
        html = render(request, "reports_export_pdf.html", {"report": data}).content.decode("utf-8")
        try:
            from weasyprint import HTML

            pdf_bytes = HTML(string=html, base_url=request.build_absolute_uri("/")).write_pdf()
        except (ImportError, OSError):
            # Windows developer environments may not have WeasyPrint's native GTK/Pango
            # libraries. Chromium is already a project dependency for visual checks.
            try:
                from playwright.sync_api import sync_playwright

                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch(headless=True)
                    try:
                        page = browser.new_page()
                        page.set_content(html, wait_until="load", timeout=15000)
                        pdf_bytes = page.pdf(format="A4", print_background=True)
                    finally:
                        browser.close()
            except Exception:
                logger.exception("No se pudo generar el PDF del reporte con WeasyPrint ni Chromium.")
                return HttpResponse("No se pudo generar el PDF en este entorno.", status=503)
        response = HttpResponse(pdf_bytes, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="reportes-actas.pdf"'
        return response
    return HttpResponse("Formato no soportado.", status=404)


@login_required
def pilot_feedback(request):
    if request.method != "POST":
        return redirect("accounts:dashboard")
    if not _pilot_active():
        messages.error(request, "El periodo de piloto no está activo.")
        return redirect("accounts:dashboard")
    message = (request.POST.get("message") or "").strip()
    category = request.POST.get("category", "ISSUE")
    if not message or len(message) > 5000:
        messages.error(request, "Escribe un comentario de hasta 5000 caracteres.")
        return redirect(request.META.get("HTTP_REFERER", "accounts:dashboard"))
    if category not in dict(PilotFeedback.CATEGORIES):
        messages.error(request, "Selecciona un tipo de comentario válido.")
        return redirect(request.META.get("HTTP_REFERER", "accounts:dashboard"))
    PilotFeedback.objects.create(user=request.user, category=category, message=message, page=request.POST.get("page", "")[:255])
    messages.success(request, "Gracias. Tu comentario quedó registrado para el equipo del piloto.")
    return redirect(request.META.get("HTTP_REFERER", "accounts:dashboard"))


def _pilot_active():
    if not settings.PILOT_MODE or not settings.PILOT_END_DATE:
        return False
    end_date = _parse_date(settings.PILOT_END_DATE)
    return bool(end_date and end_date >= timezone.localdate())
